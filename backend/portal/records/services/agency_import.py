"""Importing agencies from a spreadsheet (P20e, design 168; owner's yes, DECISIONS 37c).

A person uploads comma-separated text or an Excel workbook (.xlsx); nothing here ever opens a
file by name. The bytes are read twice: once for the preview, which says of every row whether it
is new, matches an agency already in the directory, or has a problem, and once more for the
import itself, inside one transaction, so what is written is decided against the directory as it
stands at that moment and not as it stood when the preview was drawn.

The rules:

- A row is matched by its name and jurisdiction, each folded to lower case with punctuation and
  runs of spaces dropped. A row may also match an agency by the name of one merged into it
  (P16): `directory.same_agency_ids` gives every name a kept agency is known by, and the match is
  always the kept agency, never the merged row.
- A row has a problem when it has no name, no jurisdiction (an agency cannot exist without one),
  a filing method or kind of portal the directory does not use, an email that is not an address,
  or when it names the same agency as an earlier row of the file. A row with a problem is never
  imported, in part or at all.
- A new row is created through `directory.create_agency`. A matching row is skipped unless the
  person chose to fill it in, and then only the fields that are blank on the agency are written,
  through `directory.update_agency` (one audit line each). Nothing is ever deleted or overwritten.
- With no "files by" column, a row with an email files by email and one with only a web form
  address files by portal; nothing else is guessed.
- One activity line names how many were added and how many filled in.
"""

import csv
import datetime
import io
import re
import zipfile
from dataclasses import dataclass, field
from uuid import UUID

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import connection, transaction
from openpyxl import load_workbook

from portal.core.services import activity as core_activity
from portal.records.actors import Actor
from portal.records.models import Agency
from portal.records.services import directory
from portal.records.services.common import ServiceError, require_human

# What a column can be read as, in the order the preview offers them, with the words people
# write in a header for each. A header is matched after the same folding a name gets.
IMPORT_FIELDS: dict[str, str] = {
    "name": "Name",
    "short_name": "Short name",
    "jurisdiction": "Jurisdiction",
    "records_law": "Records law",
    "filing_method": "Files by",
    "portal_kind": "Kind of portal",
    "default_request_email": "Records email",
    "portal_url": "Web form address",
    "recipient_office": "Letter goes to",
    "mailing_street": "Street",
    "mailing_city_state_zip": "City, state, ZIP",
    "notes": "Notes",
}
_HEADER_WORDS: dict[str, tuple[str, ...]] = {
    "name": ("name", "agency", "agency name", "department", "office name"),
    "short_name": ("short name", "abbreviation", "abbr", "short"),
    "jurisdiction": ("jurisdiction", "state"),
    "records_law": ("records law", "law"),
    "filing_method": ("files by", "filing method", "method", "how to file", "channel"),
    "portal_kind": ("kind of portal", "portal kind", "portal type"),
    "default_request_email": ("email", "e mail", "records email", "request email"),
    "portal_url": ("portal", "web form", "web form address", "portal url", "form", "url"),
    "recipient_office": ("letter goes to", "office", "recipient", "recipient office"),
    "mailing_street": ("street", "address", "mailing address", "street address"),
    "mailing_city_state_zip": ("city state zip", "city"),
    "notes": ("notes", "note", "comments"),
}
# The words a "files by" cell may hold, and what each means in the directory's own vocabulary.
_FILING_WORDS: dict[str, tuple[str, str | None]] = {
    "email": ("email", None),
    "e mail": ("email", None),
    "portal": ("portal", None),
    "web form": ("portal", "Web form"),
    "webform": ("portal", "Web form"),
    "govqa": ("portal", "GovQA"),
    "nextrequest": ("portal", "NextRequest"),
    "next request": ("portal", "NextRequest"),
    "mail": ("mail", None),
    "post": ("mail", None),
    "postal": ("mail", None),
    "postal mail": ("mail", None),
}
_KINDS = {"".join(kind.casefold().split()): kind for kind in directory.PORTAL_KINDS}

# A spreadsheet of agencies is small; these stop a mistaken upload, not a real directory.
MAX_BYTES = 2_000_000
MAX_ROWS = 5_000
# A workbook is a ZIP, and a small one can unpack to gigabytes; this is checked before it is read.
MAX_WORKBOOK_UNPACKED_BYTES = 50 * 1024 * 1024

STATUSES = ("new", "match", "problem")
DECISIONS = ("skip", "update")


def fold(text: str) -> str:
    """Lower case, punctuation and runs of spaces dropped: "U.S. Customs" is "u s customs"."""
    return " ".join(re.sub(r"[^\w\s]", " ", text.casefold()).split())


@dataclass(frozen=True)
class ImportRow:
    """One row of the file as the import would treat it."""

    number: int  # the row's place in the spreadsheet; the header is usually row 1
    values: dict[str, str]
    status: str  # one of STATUSES
    problems: tuple[str, ...] = ()
    match_id: UUID | None = None
    match_name: str | None = None
    # The agency's blank fields this row would fill in, if the person chooses to.
    fills: tuple[str, ...] = ()


@dataclass(frozen=True)
class ImportPreview:
    headers: list[str]
    mapping: dict[str, str | None]
    rows: list[ImportRow] = field(default_factory=list)

    def count(self, status: str) -> int:
        return sum(1 for row in self.rows if row.status == status)


@dataclass(frozen=True)
class ImportResult:
    created: int
    updated: int
    skipped: int
    problems: int
    summary: str


def _rows_from_text(data: bytes) -> list[list[str]]:
    """Every row of comma-separated text, each cell stripped. Commas, unless the file plainly uses
    semicolons or tabs (a spreadsheet saved in a locale that writes decimals with commas)."""
    text = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ServiceError(
            "That file is neither comma-separated text nor an Excel workbook (.xlsx)."
        )
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t"
        )
    except csv.Error:
        dialect = csv.excel
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _cell_text(value: object) -> str:
    """A workbook cell as the text a comma-separated file would hold: a blank is empty, a whole
    number has no ".0" (a ZIP code stored as 92093.0 is 92093), a date is its ISO date."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value).strip()


def _rows_from_workbook(data: bytes) -> list[list[str]]:
    """The rows of the first worksheet of an .xlsx workbook, values only (a formula gives the
    answer it last showed, never the formula). Reading stops once there are more rows than an
    agency list should have, so a sheet of a million rows is not read to the end."""
    unreadable = ServiceError(
        "That workbook could not be read: save it again as .xlsx or as comma-separated text."
    )
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unpacked = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile:
        raise unreadable from None
    if unpacked > MAX_WORKBOOK_UNPACKED_BYTES:
        raise ServiceError("That workbook is too big for an agency list.")
    rows: list[list[str]] = []
    filled = 0
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for raw in workbook.worksheets[0].iter_rows(values_only=True):
                row = [_cell_text(value) for value in raw]
                rows.append(row)
                filled += any(row)
                if filled > MAX_ROWS + 1:
                    break
        finally:
            workbook.close()
    except Exception:  # openpyxl raises many kinds of error for a damaged or odd workbook
        raise unreadable from None
    return rows


def _table(data: bytes) -> list[tuple[int, list[str]]]:
    """The non-empty rows with their place in the sheet (the first row is 1), each cell
    stripped, from comma-separated text or an .xlsx workbook. A workbook is told by its content
    (a ZIP begins "PK"), never by the file's name."""
    if len(data) > MAX_BYTES:
        raise ServiceError("That file is too big for an agency list.")
    is_workbook = data[:4] == b"PK\x03\x04"
    raw = _rows_from_workbook(data) if is_workbook else _rows_from_text(data)
    rows = [(number, row) for number, row in enumerate(raw, start=1) if any(row)]
    if len(rows) > MAX_ROWS + 1:
        raise ServiceError("That file has more rows than an agency list should.")
    return rows


def _headers(row: list[str]) -> list[str]:
    """The header row, a blank header named by its place and a repeated one numbered."""
    seen: dict[str, int] = {}
    out = []
    for index, cell in enumerate(row, start=1):
        name = cell or f"Column {index}"
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return out


def suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    """Each header's field, by the words people write; the first header wins a field."""
    taken: set[str] = set()
    mapping: dict[str, str | None] = {}
    for header in headers:
        folded = fold(header)
        found = next(
            (
                key
                for key, words in _HEADER_WORDS.items()
                if key not in taken and (folded in words or folded == fold(IMPORT_FIELDS[key]))
            ),
            None,
        )
        if found:
            taken.add(found)
        mapping[header] = found
    return mapping


def _check_mapping(headers: list[str], mapping: dict[str, str | None]) -> dict[str, str | None]:
    out: dict[str, str | None] = {}
    used: set[str] = set()
    for header in headers:
        target = mapping.get(header) or None
        if target is not None:
            if target not in IMPORT_FIELDS:
                raise ServiceError(f"There is no field called {target!r} to import into.")
            if target in used:
                raise ServiceError(f"Two columns are both read as {IMPORT_FIELDS[target]}.")
            used.add(target)
        out[header] = target
    return out


def _parse_filing(values: dict[str, str], problems: list[str]) -> None:
    """Turn the "files by" and "kind of portal" cells into the directory's words, in place."""
    method_cell = values.pop("filing_method", "")
    kind_cell = values.pop("portal_kind", "")
    method: str | None = None
    kind: str | None = None
    if method_cell:
        meaning = _FILING_WORDS.get(fold(method_cell))
        if meaning is None:
            problems.append(f"Files by “{method_cell}” is not a way the directory files")
        else:
            method, kind = meaning
    if kind_cell:
        named = _KINDS.get("".join(kind_cell.casefold().split()))
        if named is None:
            problems.append(f"“{kind_cell}” is not a kind of portal the directory knows")
        elif method not in (None, "portal"):
            problems.append("A kind of portal is given for an agency that does not file by portal")
        else:
            method, kind = "portal", named
    if method is None and not method_cell:
        # Nothing said how: an email files by email, a web form address alone by portal.
        if values.get("default_request_email"):
            method = "email"
        elif values.get("portal_url"):
            method = "portal"
    if method:
        values["filing_method"] = method
    if kind:
        values["portal_kind"] = kind


def _directory_index() -> dict[tuple[str, str], Agency]:
    """Every name an agency in the directory goes by, with its jurisdiction, to the agency kept.

    An agency merged into another is the same agency (P16): its name and jurisdiction lead to the
    one kept, found through `directory.same_agency_ids`, which is the read-side answer to "the
    same agency"."""
    kept = {a.pk: a for a in Agency.objects.filter(merged_into__isnull=True)}
    index: dict[tuple[str, str], Agency] = {}
    for agency in kept.values():
        index.setdefault((fold(agency.name), fold(agency.jurisdiction)), agency)
    with_aliases = set(
        Agency.objects.filter(merged_into__isnull=False).values_list("merged_into_id", flat=True)
    )
    names = dict(
        (pk, (name, jurisdiction))
        for pk, name, jurisdiction in Agency.objects.values_list("pk", "name", "jurisdiction")
    )
    for kept_id in with_aliases:
        if kept_id not in kept:
            continue
        for same_id in directory.same_agency_ids(kept_id):
            name, jurisdiction = names[same_id]
            index.setdefault((fold(name), fold(jurisdiction)), kept[kept_id])
    return index


def _blank_fills(agency: Agency, values: dict[str, str]) -> tuple[str, ...]:
    """The fields this row has and the agency leaves blank. Never the name or jurisdiction (they
    are what matched), and the two filing fields only as the pair the directory accepts."""
    fills = [
        key
        for key in IMPORT_FIELDS
        if key not in ("name", "jurisdiction", "filing_method", "portal_kind")
        and values.get(key)
        and not getattr(agency, key)
    ]
    method = values.get("filing_method")
    if method and not agency.filing_method:
        fills.append("filing_method")
        if values.get("portal_kind"):
            fills.append("portal_kind")
    elif values.get("portal_kind") and agency.filing_method == "portal" and not agency.portal_kind:
        fills.append("portal_kind")
    return tuple(fills)


def preview(data: bytes, mapping: dict[str, str | None] | None = None) -> ImportPreview:
    """What the import of these bytes would do, row by row, against the directory now. Read only."""
    table = _table(data)
    if not table:
        raise ServiceError("That file is empty.")
    headers = _headers(table[0][1])
    chosen = _check_mapping(headers, suggest_mapping(headers) if mapping is None else mapping)
    index = _directory_index()
    seen: dict[tuple[str, str], int] = {}
    rows: list[ImportRow] = []
    for number, cells in table[1:]:
        values = {
            target: cells[i]
            for i, header in enumerate(headers)
            if (target := chosen.get(header)) and i < len(cells) and cells[i]
        }
        problems: list[str] = []
        if not values.get("name"):
            problems.append("No name")
        if not values.get("jurisdiction"):
            problems.append("No jurisdiction")
        email = values.get("default_request_email")
        if email:
            try:
                validate_email(email)
            except ValidationError:
                problems.append(f"“{email}” is not an email address")
        _parse_filing(values, problems)
        key = (fold(values.get("name", "")), fold(values.get("jurisdiction", "")))
        if not problems and key in seen:
            problems.append(f"The same agency as row {seen[key]}")
        if problems:
            rows.append(ImportRow(number, values, "problem", tuple(problems)))
            continue
        seen[key] = number
        match = index.get(key)
        if match is None:
            rows.append(ImportRow(number, values, "new"))
        else:
            rows.append(
                ImportRow(
                    number,
                    values,
                    "match",
                    match_id=match.pk,
                    match_name=match.name,
                    fills=_blank_fills(match, values),
                )
            )
    return ImportPreview(headers=headers, mapping=chosen, rows=rows)


@transaction.atomic
def import_agencies(
    actor: Actor,
    data: bytes,
    *,
    filename: str,
    mapping: dict[str, str | None] | None = None,
    decisions: dict[int, str] | None = None,
) -> ImportResult:
    """Create the new rows and fill in the matching ones the person chose, all or nothing.

    `decisions` maps a matching row's number to "update" (fill in its blank fields) or "skip";
    a matching row with no decision is skipped. The rows are classified again here, under the
    transaction, so an agency added since the preview is matched rather than duplicated."""
    require_human(actor, "Importing agencies")
    decisions = decisions or {}
    unknown = sorted(set(decisions.values()) - set(DECISIONS))
    if unknown:
        raise ServiceError(f"A matching row is either skipped or filled in, not {unknown[0]!r}.")
    # One import at a time, so two people importing the same file cannot both create a row.
    with connection.cursor() as cursor:
        cursor.execute("SELECT pg_advisory_xact_lock(hashtext(%s))", ["agency-import"])
    found = preview(data, mapping)
    name = (filename or "the file").strip()
    created = updated = skipped = 0
    created_ids: list[str] = []
    for row in found.rows:
        if row.status == "new":
            made = directory.create_agency(
                actor,
                name=row.values["name"],
                jurisdiction=row.values["jurisdiction"],
                imported_from=name,
                **{k: v for k, v in row.values.items() if k not in ("name", "jurisdiction")},
            )
            created_ids.append(str(made.pk))
            created += 1
        elif row.status == "match" and decisions.get(row.number) == "update" and row.fills:
            assert row.match_id is not None
            directory.update_agency(
                actor,
                row.match_id,
                f"Filled in from an import of {name}.",
                **{key: row.values[key] for key in row.fills},
            )
            updated += 1
        elif row.status == "match":
            skipped += 1
    if created + updated == 0:
        raise ServiceError("Nothing in this file would be added or filled in.")
    added = f"{created} {'agency' if created == 1 else 'agencies'} added"
    summary = f"Imported agencies from {name}: {added}, {updated} filled in"
    core_activity.record(
        actor,
        "agencies_imported",
        summary,
        subject_type="agencies",
        data={
            "filename": name,
            "created": created,
            "agency_ids": created_ids,
            "updated": updated,
            "skipped": skipped,
            "problems": found.count("problem"),
        },
    )
    return ImportResult(created, updated, skipped, found.count("problem"), summary)
